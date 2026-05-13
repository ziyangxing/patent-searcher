import os
import uuid
from fastapi import APIRouter, UploadFile, File, HTTPException, Depends
from sqlalchemy.ext.asyncio import AsyncSession
from app.core.config import settings
from app.db.session import get_db
from app.schemas.patent import UploadedDocResponse
from app.models.patent import UploadedDocument

router = APIRouter()


@router.post("/patent", response_model=UploadedDocResponse)
async def upload_patent(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
):
    """上传专利文档（PDF/TXT），提取文本内容"""
    if not file.filename:
        raise HTTPException(status_code=400, detail="No file provided")

    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in [".pdf", ".txt", ".docx", ".doc"]:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file type: {ext}. Supported: PDF, TXT, DOCX",
        )

    content = await file.read()
    if len(content) > settings.MAX_UPLOAD_SIZE:
        raise HTTPException(status_code=400, detail="File too large. Max 50MB")

    os.makedirs(settings.UPLOAD_DIR, exist_ok=True)
    file_id = str(uuid.uuid4())
    save_name = f"{file_id}{ext}"
    save_path = os.path.join(settings.UPLOAD_DIR, save_name)

    with open(save_path, "wb") as f:
        f.write(content)

    extracted_text = ""
    if ext == ".txt":
        extracted_text = content.decode("utf-8", errors="ignore")
    elif ext == ".pdf":
        try:
            import fitz

            doc = fitz.open(stream=content, filetype="pdf")
            for page in doc:
                extracted_text += page.get_text()
            doc.close()
        except Exception as e:
            raise HTTPException(
                status_code=500, detail=f"Failed to parse PDF: {str(e)}"
            )
    elif ext in [".docx", ".doc"]:
        try:
            from docx import Document

            doc = Document(save_path)
            extracted_text = "\n".join(
                p.text for p in doc.paragraphs
            )
        except Exception as e:
            raise HTTPException(
                status_code=500, detail=f"Failed to parse DOCX: {str(e)}"
            )

    record = UploadedDocument(
        user_id="default",
        original_name=file.filename,
        file_path=save_path,
        file_type=ext.replace(".", ""),
        extracted_text=extracted_text,
        status="processed",
    )
    db.add(record)
    await db.commit()
    await db.refresh(record)

    return UploadedDocResponse(
        id=record.id,
        original_name=record.original_name,
        file_type=record.file_type,
        extracted_text=extracted_text[:2000],
        patent_number=record.patent_number,
        status=record.status,
        created_at=record.created_at,
    )
